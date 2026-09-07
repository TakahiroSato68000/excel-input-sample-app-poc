import json
import shutil
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from services.excel_service import create_template_if_missing
from services.storage_service import ensure_runtime_directories


class AppTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp(prefix="excel-poc-"))
        ensure_runtime_directories(self.temp_dir)
        self._write_roster()
        (self.temp_dir / "template").mkdir(parents=True, exist_ok=True)
        shutil.copy2(
            Path.cwd() / "template" / "template.xlsx",
            self.temp_dir / "template" / "template.xlsx",
        )
        create_template_if_missing(self.temp_dir)
        self._patched_app, self._restore_patches = create_app_for(self.temp_dir)
        self.client = TestClient(self._patched_app)

    def tearDown(self) -> None:
        self._restore_patches()
        for child in sorted(self.temp_dir.rglob("*"), reverse=True):
            if child.is_file():
                child.unlink()
            elif child.is_dir():
                child.rmdir()

    def _write_roster(self) -> None:
        roster = [
            {"employee_number": "001", "employee_name": "山田 太郎"},
            {"employee_number": "002", "employee_name": "鈴木 花子"},
        ]
        (self.temp_dir / "roster.json").write_text(json.dumps(roster, ensure_ascii=False), encoding="utf-8")

    def test_save_and_list_entries(self) -> None:
        response = self.client.post(
            "/api/entries",
            json={
                "date": "2026-09-07",
                "employee_name": "山田 太郎",
                "employee_number": "",
                "reason": "対応",
                "start_time": "09:00",
                "end_time": "17:00",
                "remarks": "",
            },
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["item"]["employee_number"], "001")

        response = self.client.get("/api/entries", params={"date": "2026-09-07"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["items"]), 1)

    def test_duplicate_entry_replaces_existing_row(self) -> None:
        self.client.post(
            "/api/entries",
            json={
                "date": "2026-09-07",
                "employee_name": "山田 太郎",
                "employee_number": "001",
                "reason": "最初",
                "start_time": "09:00",
                "end_time": "17:00",
                "remarks": "",
            },
        )
        self.client.post(
            "/api/entries",
            json={
                "date": "2026-09-07",
                "employee_name": "山田 太郎",
                "employee_number": "001",
                "reason": "更新",
                "start_time": "10:00",
                "end_time": "18:00",
                "remarks": "note",
            },
        )
        records = self.client.get("/api/entries", params={"date": "2026-09-07"}).json()["items"]
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["reason"], "更新")

    def test_confirm_generates_excel(self) -> None:
        self.client.post(
            "/api/entries",
            json={
                "date": "2026-09-07",
                "employee_name": "山田 太郎",
                "employee_number": "001",
                "reason": "対応",
                "start_time": "09:00",
                "end_time": "17:00",
                "remarks": "",
            },
        )
        response = self.client.post("/api/confirm", json={"date": "2026-09-07"})
        self.assertEqual(response.status_code, 200)

        output_file = self.temp_dir / "2026-09-07.xlsx"
        self.assertTrue(output_file.exists())
        workbook = load_workbook(output_file)
        self.assertEqual(workbook.sheetnames, ["Print", "Data"])
        self.assertIsInstance(workbook["Print"]["C4"].value, str)
        self.assertTrue(workbook["Print"]["C4"].value.startswith("=IF("))
        self.assertIn("Data!D2", workbook["Print"]["F4"].value)
        self.assertIn("Data!F2", workbook["Print"]["H4"].value)
        self.assertGreater(sum(1 for row in workbook["Print"].iter_rows() for cell in row if cell.has_style), 0)

        data_sheet = workbook["Data"]
        self.assertEqual(data_sheet["A2"].value, "2026-09-07")
        self.assertEqual(data_sheet["C2"].value, "山田 太郎")
        self.assertEqual(data_sheet["D2"].value, "09:00")
        self.assertEqual(data_sheet["E2"].value, "17:00")
        self.assertEqual(data_sheet["F2"].value, "対応")


def create_app_for(base_dir: Path):
    from fastapi import FastAPI
    from fastapi.staticfiles import StaticFiles
    from routes import router
    from services.storage_service import ensure_runtime_directories

    ensure_runtime_directories(base_dir)
    app = FastAPI(title="test-app")
    app.mount("/static", StaticFiles(directory=base_dir / "static"), name="static")
    app.include_router(router)

    # Override module-level paths used by the router/services for the test instance,
    # keeping the originals so they can be restored once the test completes.
    import routes as routes_module
    import services.storage_service as storage_module
    import services.excel_service as excel_module
    import services.roster_service as roster_module

    originals = {
        (routes_module, "BASE_DIR"): routes_module.BASE_DIR,
        (routes_module, "INDEX_FILE"): routes_module.INDEX_FILE,
        (storage_module, "BASE_DIR"): storage_module.BASE_DIR,
        (excel_module, "template_path"): excel_module.template_path,
        (excel_module, "backup_dir"): excel_module.backup_dir,
        (roster_module, "roster_path"): roster_module.roster_path,
    }

    routes_module.BASE_DIR = base_dir
    routes_module.INDEX_FILE = base_dir / "templates" / "index.html"
    storage_module.BASE_DIR = base_dir
    excel_module.template_path = lambda _base_dir: base_dir / "template" / "template.xlsx"
    excel_module.backup_dir = lambda _base_dir: base_dir / "backup"
    roster_module.roster_path = lambda _base_dir: base_dir / "roster.json"

    def restore() -> None:
        for (module, attr), value in originals.items():
            setattr(module, attr, value)

    return app, restore
