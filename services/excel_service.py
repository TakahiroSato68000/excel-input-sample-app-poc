from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from services.errors import ExcelGenerationError, StorageError
from services.storage_service import (
    data_file_path,
    load_records_for_date,
    output_file_path,
)


HEADER_ROW = [
    "日付",
    "社員番号",
    "氏名",
    "理由",
    "開始時刻",
    "終了時刻",
    "備考",
]


def template_path(base_dir: Path) -> Path:
    return base_dir / "template.xlsx"


def backup_dir(base_dir: Path) -> Path:
    return base_dir / "backup"


def ensure_output_backup(base_dir: Path, output_path: Path) -> None:
    if not output_path.exists():
        return
    target_dir = backup_dir(base_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
    backup_path = target_dir / f"{output_path.stem}_{timestamp}{output_path.suffix}"
    # Guard against filename collisions (e.g. repeated confirms within the same microsecond).
    suffix_counter = 1
    while backup_path.exists():
        backup_path = target_dir / f"{output_path.stem}_{timestamp}_{suffix_counter}{output_path.suffix}"
        suffix_counter += 1
    try:
        shutil.copy2(output_path, backup_path)
    except OSError as exc:
        raise StorageError("backup failed") from exc


def fill_data_sheet(workbook: Workbook, records: list[dict[str, str]]) -> None:
    sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.active
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    if sheet.max_row == 1 and sheet["A1"].value is None:
        for index, title in enumerate(HEADER_ROW, start=1):
            sheet.cell(row=1, column=index, value=title)
    for row_index, record in enumerate(records, start=2):
        sheet.cell(row=row_index, column=1, value=record["date"])
        sheet.cell(row=row_index, column=2, value=record["employee_number"])
        sheet.cell(row=row_index, column=3, value=record["employee_name"])
        sheet.cell(row=row_index, column=4, value=record["reason"])
        sheet.cell(row=row_index, column=5, value=record["start_time"])
        sheet.cell(row=row_index, column=6, value=record["end_time"])
        sheet.cell(row=row_index, column=7, value=record["remarks"])

    print_sheet = workbook["Print"] if "Print" in workbook.sheetnames else workbook.create_sheet("Print")
    print_sheet["A1"] = "印刷用シート"
    print_sheet["A2"] = "Data シートを参照して出力しています。"


def create_template_if_missing(base_dir: Path) -> None:
    path = template_path(base_dir)
    if path.exists():
        return
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Data"
    for index, title in enumerate(HEADER_ROW, start=1):
        data_sheet.cell(row=1, column=index, value=title)
    print_sheet = workbook.create_sheet("Print")
    print_sheet["A1"] = "印刷用シート"
    print_sheet["A2"] = "Data シートを参照して出力しています。"
    workbook.save(path)


def confirm_records(base_dir: Path, target_date: str) -> dict[str, Any]:
    data_path = data_file_path(base_dir, target_date)
    if not data_path.exists():
        raise ValueError("no saved data for the selected date")

    records = load_records_for_date(base_dir, target_date)
    output_path = output_file_path(base_dir, target_date)
    ensure_output_backup(base_dir, output_path)

    template = template_path(base_dir)
    if not template.exists():
        raise ExcelGenerationError("template.xlsx not found")

    workbook = load_workbook(template)
    fill_data_sheet(workbook, records)
    workbook.save(output_path)
    return {
        "message": "Excel を生成しました。",
        "output_file": output_path.name,
        "record_count": len(records),
    }

